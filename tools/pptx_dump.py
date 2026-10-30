import sys
from pptx import Presentation

def walk(shapes, n):
    for s in shapes:
        if s.shape_type == 6:
            walk(s.shapes, n)
            continue
        if s.has_table if hasattr(s, "has_table") else False:
            for r, row in enumerate(s.table.rows):
                for c, cell in enumerate(row.cells):
                    t = cell.text.replace("\n", "⏎")[:40]
                    if t:
                        print(f"S{n}|{s.shape_id}|TABLE r{r}c{c}|{s.name}|{t}")
        elif getattr(s, "has_chart", False) and s.has_chart:
            ch = s.chart
            cats = list(ch.plots[0].categories)
            for se in ch.plots[0].series:
                print(f"S{n}|{s.shape_id}|CHART|{s.name}|{cats}={list(se.values)}")
        elif s.shape_type == 13:
            print(f"S{n}|{s.shape_id}|PICTURE|{s.name}|")
        elif s.has_text_frame and s.text_frame.text.strip():
            t = s.text_frame.text.replace("\n", "⏎")[:40]
            print(f"S{n}|{s.shape_id}|TEXT|{s.name}|{t}")

prs = Presentation(sys.argv[1])
wanted = {int(x) for x in sys.argv[2:]}
for i, sl in enumerate(prs.slides, 1):
    if not wanted or i in wanted:
        walk(sl.shapes, i)
